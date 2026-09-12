import os
import json
import re
import zipfile
import xml.etree.ElementTree as ET

def parse_page_range_str(pages_str):
    """
    Parses slide page numbers or ranges (e.g. '1', '1,3-5', '2,4,6-8')
    into a set of 1-based integer slide numbers.
    Returns None if pages_str is empty/None (meaning all slides).
    """
    if not pages_str:
        return None
    page_numbers = set()
    parts = str(pages_str).split(",")
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            subparts = part.split("-")
            if len(subparts) == 2 and subparts[0].isdigit() and subparts[1].isdigit():
                start_p = int(subparts[0])
                end_p = int(subparts[1])
                for p in range(min(start_p, end_p), max(start_p, end_p) + 1):
                    page_numbers.add(p)
        elif part.isdigit():
            page_numbers.add(int(part))

    return page_numbers if page_numbers else None


def extract_pptx_images(pptx_path, target_image_dir, page_numbers=None):
    """
    Extracts images from specified slide pages of a PPTX file.
    Supports python-pptx if installed, with a zero-dependency ZipFile XML fallback.
    Returns a list of image dicts: [{'id': ..., 'path': ..., 'title': ..., 'slide_num': ...}]
    """
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

    os.makedirs(target_image_dir, exist_ok=True)
    pptx_base = os.path.splitext(os.path.basename(pptx_path))[0]
    extracted_images = []

    # Try python-pptx first
    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        for slide_idx, slide in enumerate(prs.slides, start=1):
            if page_numbers and slide_idx not in page_numbers:
                continue
            img_count = 1
            for shape in slide.shapes:
                if hasattr(shape, "image") and shape.image:
                    blob = shape.image.blob
                    ext = f".{shape.image.ext}" if shape.image.ext else ".png"
                    img_id = f"{pptx_base}_slide_{slide_idx}_img_{img_count}"
                    filename = f"{img_id}{ext}"
                    out_path = os.path.join(target_image_dir, filename)
                    with open(out_path, "wb") as f:
                        f.write(blob)

                    extracted_images.append({
                        "id": img_id,
                        "path": os.path.relpath(out_path, start=os.path.dirname(os.path.abspath(target_image_dir)) or "."),
                        "abs_path": out_path,
                        "title": f"{pptx_base} Slide {slide_idx}",
                        "slide_num": slide_idx
                    })
                    img_count += 1
        if extracted_images:
            return extracted_images
    except Exception:
        pass

    # Zero-dependency ZipFile fallback
    with zipfile.ZipFile(pptx_path, "r") as z:
        file_list = z.namelist()

        # Parse ppt/presentation.xml for slide ordering
        slide_rids = []
        if "ppt/presentation.xml" in file_list:
            try:
                tree = ET.fromstring(z.read("ppt/presentation.xml"))
                ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
                      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
                for sld in tree.findall(".//p:sldId", ns):
                    r_id = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                    if r_id:
                        slide_rids.append(r_id)
            except Exception:
                pass

        # Map rIds to slide paths using ppt/_rels/presentation.xml.rels
        rel_map = {}
        if "ppt/_rels/presentation.xml.rels" in file_list:
            try:
                tree = ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
                for rel in tree.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                    r_id = rel.get("Id")
                    target = rel.get("Target")
                    if r_id and target and "slides/slide" in target:
                        rel_map[r_id] = target.replace("slides/", "")
            except Exception:
                pass

        ordered_slides = []
        for r_id in slide_rids:
            if r_id in rel_map:
                ordered_slides.append(rel_map[r_id])

        if not ordered_slides:
            # Fallback: sort slide XML files numerically
            slide_files = [f for f in file_list if re.match(r"^ppt/slides/slide\d+\.xml$", f)]
            slide_files.sort(key=lambda x: int(re.search(r"\d+", x).group()))
            ordered_slides = [os.path.basename(f) for f in slide_files]

        for slide_idx, slide_file in enumerate(ordered_slides, start=1):
            if page_numbers and slide_idx not in page_numbers:
                continue

            rel_file = f"ppt/slides/_rels/{slide_file}.rels"
            if rel_file not in file_list:
                continue

            try:
                rel_tree = ET.fromstring(z.read(rel_file))
                img_targets = []
                for rel in rel_tree.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                    target = rel.get("Target", "")
                    if "media/" in target:
                        media_path = "ppt/" + target.lstrip("../")
                        if media_path in file_list:
                            img_targets.append(media_path)

                img_targets.sort()
                for img_count, media_path in enumerate(img_targets, start=1):
                    ext = os.path.splitext(media_path)[1] or ".png"
                    img_id = f"{pptx_base}_slide_{slide_idx}_img_{img_count}"
                    filename = f"{img_id}{ext}"
                    out_path = os.path.join(target_image_dir, filename)

                    with open(out_path, "wb") as out_f:
                        out_f.write(z.read(media_path))

                    extracted_images.append({
                        "id": img_id,
                        "path": filename,
                        "abs_path": out_path,
                        "title": f"{pptx_base} Slide {slide_idx}",
                        "slide_num": slide_idx
                    })
            except Exception:
                pass

    return extracted_images


def add_pptx_to_record(pptx_path, record_json_path="record.json", page_numbers=None, title=None, grid="1x1", image_dir="images", force=False):
    """
    Extracts images from PPTX slides, registers them in record.json,
    and appends an image section formatted with the specified grid size (e.g. '2x2', '3x5', '1x2', etc.).
    """
    if not os.path.exists(record_json_path):
        raise FileNotFoundError(f"Target document JSON file not found: {record_json_path}")

    with open(record_json_path, "r", encoding="utf-8") as f:
        try:
            record_data = json.load(f)
        except Exception as e:
            raise ValueError(f"Failed to parse target JSON file {record_json_path}: {e}")

    parsed_pages = parse_page_range_str(page_numbers) if isinstance(page_numbers, str) else page_numbers
    doc_dir = os.path.dirname(os.path.abspath(record_json_path))
    target_img_dir = os.path.join(doc_dir, image_dir) if not os.path.isabs(image_dir) else image_dir

    extracted = extract_pptx_images(pptx_path, target_img_dir, page_numbers=parsed_pages)
    if not extracted:
        raise ValueError(f"No images could be extracted from PPTX file '{pptx_path}'.")

    # Update global images registry in record.json
    existing_images = record_data.setdefault("images", [])
    existing_ids = {img.get("id"): img for img in existing_images if isinstance(img, dict)}

    new_img_ids = []
    for count, item in enumerate(extracted, start=1):
        img_id = item["id"]
        rel_path = os.path.join(image_dir, os.path.basename(item["path"])) if image_dir else item["path"]
        img_entry = {
            "id": img_id,
            "path": rel_path,
            "title": str(count)
        }
        if img_id not in existing_ids or force:
            if img_id in existing_ids:
                existing_images.remove(existing_ids[img_id])
            existing_images.append(img_entry)
            existing_ids[img_id] = img_entry
        new_img_ids.append(img_id)

    # Determine items per page based on grid specification (e.g., '2x2', '3x5')
    grid_str = str(grid).strip()
    grid_match = re.match(r"^(\d+)x(\d+)$", grid_str.lower())
    if grid_match:
        per_page = int(grid_match.group(1)) * int(grid_match.group(2))
    else:
        per_page = 1
        grid_str = "1x1"

    # Build image subsections
    subsections = []
    page_counter = 1
    for i in range(0, len(new_img_ids), per_page):
        chunk = new_img_ids[i:i + per_page]
        subsections.append({
            "title": f"Page {page_counter}",
            "layout": {
                "type": grid_str,
                "elements": chunk
            }
        })
        page_counter += 1

    pptx_base = os.path.splitext(os.path.basename(pptx_path))[0]
    sec_title = title if title else f"PPTX Import ({pptx_base})"

    new_section = {
        "type": "image",
        "title": sec_title,
        "new_page": True,
        "subsections": subsections
    }

    experiments = record_data.setdefault("experiments", [])
    if not experiments:
        experiments.append({
            "number": 1,
            "title": "Lab Experiment",
            "sections": []
        })

    experiments[0].setdefault("sections", []).append(new_section)

    with open(record_json_path, "w", encoding="utf-8") as f:
        json.dump(record_data, f, indent=4)

    return new_section, len(extracted)
